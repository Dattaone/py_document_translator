import os

from docx import Document
from domain.entities.text_unit import TextUnit, TextUnitType
from infrastructure.config.paths import OUTPUT_DIR
from infrastructure.config.paths import TRANSLATIONS_DIR

import logging

logger = logging.getLogger(__name__)


class DocxBuilder:
    def build(
            self, 
            translated_units: list[TextUnit], 
            path: str):
        try:
            file_path = os.path.join(OUTPUT_DIR, path)

            doc = Document()

            for unit in translated_units:
                if unit.type == TextUnitType.HEADING:
                    doc.add_heading(f"{unit.text}\n", level=1)
                    continue
                
                p = doc.add_paragraph(style=unit.metadata["style"])
                for run in unit.metadata["runs"]:
                    r = p.add_run(run["text"])
                    r.bold = run["bold"]
                    r.italic = run["italic"]
                    r.underline = run["underline"]
                p.add_run("\n")

            doc.save(f"{TRANSLATIONS_DIR}/traducido_{path}")

            logger.info(f"✅ Documento guardado como {TRANSLATIONS_DIR}/traducido_{path}")
            return True



            """ 
            document = Document(file_path)
            paragraph_units = document.paragraphs
            for index, paragraph_unit in enumerate(paragraph_units):
                paragraph_unit.text = translated_units[index].text
                
            
            document.save(f"{TRANSLATIONS_DIR}/traducido_{path}")
            logger.info("Documento creado") 
            """
        except Exception as e:
            logger.exception(f"Error al crear documento {e}")